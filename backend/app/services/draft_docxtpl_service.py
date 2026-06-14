"""docxtpl: рендер Word-шаблонов, preview HTML для редактора."""
from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from typing import Dict, List, Tuple

logger = logging.getLogger(__name__)

_DOCXTPL_VAR_RE = re.compile(r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*)\s*\}\}")


def scan_docxtpl_occurrences(text: str) -> List[Tuple[str, str, int, int]]:
    """
    Найти все {{ variable }} в тексте.

    Returns:
        [(variable, marker, start, end), ...]
    """
    out: List[Tuple[str, str, int, int]] = []
    for m in _DOCXTPL_VAR_RE.finditer(text or ""):
        var = m.group(1)
        out.append((var, m.group(0), m.start(), m.end()))
    return out


def docx_to_plain_text(docx_path: Path) -> str:
    """Извлечь текст из DOCX (параграфы + ячейки таблиц)."""
    from docx import Document

    doc = Document(str(docx_path))
    parts: List[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def docx_to_preview_html(docx_path: Path) -> str:
    """DOCX → HTML для редактора (сохраняет {{ variable }})."""
    from docx import Document

    doc = Document(str(docx_path))
    blocks: List[str] = []
    for para in doc.paragraphs:
        text = para.text
        if not text.strip():
            blocks.append("<p>&nbsp;</p>")
            continue
        escaped = _escape_html(text)
        blocks.append(f"<p>{escaped}</p>")
    for table in doc.tables:
        rows_html: List[str] = []
        for row in table.rows:
            cells = "".join(f"<td>{_escape_html(c.text)}</td>" for c in row.cells)
            rows_html.append(f"<tr>{cells}</tr>")
        blocks.append(f"<table>{''.join(rows_html)}</table>")
    return "\n".join(blocks) if blocks else "<p></p>"


def _escape_html(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def render_docxtpl(docx_path: Path, context: Dict[str, str]) -> bytes:
    """Рендер DOCX через docxtpl."""
    from docxtpl import DocxTemplate

    tpl = DocxTemplate(str(docx_path))
    tpl.render({k: (v or "") for k, v in context.items()})
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue()


def rendered_docx_to_html(docx_bytes: bytes) -> str:
    """Заполненный DOCX → HTML для редактора."""
    import tempfile

    from docx import Document

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(docx_bytes)
        tmp_path = Path(tmp.name)
    try:
        doc = Document(str(tmp_path))
        blocks: List[str] = []
        for para in doc.paragraphs:
            text = para.text
            if not text.strip():
                blocks.append("<p>&nbsp;</p>")
                continue
            blocks.append(f"<p>{_escape_html(text)}</p>")
        return "\n".join(blocks) if blocks else "<p></p>"
    finally:
        tmp_path.unlink(missing_ok=True)


def build_docx_from_text(body: str, out_path: Path) -> None:
    """Создать DOCX из текста (по одному параграфу на строку)."""
    from docx import Document

    doc = Document()
    for line in body.splitlines():
        doc.add_paragraph(line)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    logger.info("[draft_docxtpl] created %s", out_path)
