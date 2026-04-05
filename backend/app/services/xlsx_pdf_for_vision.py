"""
Конвертация XLSX → PDF (таблицы) для последующего OCR в Yandex Vision.
Нужен TTF с кириллицей в системе (DejaVu, Arial Unicode и т.п.).
"""
from __future__ import annotations

import io
import logging
import os
import re
from typing import List, Optional

from app.config import config

logger = logging.getLogger(__name__)

_ui_font_name: Optional[str] = None


def _find_cyrillic_ttf() -> Optional[str]:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "C:\\\\Windows\\\\Fonts\\\\arial.ttf",
        "C:\\\\Windows\\\\Fonts\\\\arialuni.ttf",
    ]
    extra = os.getenv("XLSX_PDF_FONT_TTF", "").strip()
    if extra and os.path.isfile(extra):
        candidates.insert(0, extra)
    for path in candidates:
        if path and os.path.isfile(path):
            return path
    return None


def _ensure_ui_font() -> str:
    global _ui_font_name
    if _ui_font_name:
        return _ui_font_name
    path = _find_cyrillic_ttf()
    if not path:
        raise RuntimeError(
            "Не найден TTF-шрифт с кириллицей для PDF из Excel. "
            "Установите пакет fonts-dejavu-core (Linux) или задайте XLSX_PDF_FONT_TTF=/path/to/font.ttf"
        )
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("XlsxVisionFont", path))
    _ui_font_name = "XlsxVisionFont"
    logger.info("XLSX→PDF: зарегистрирован шрифт %s", path)
    return _ui_font_name


def _cell_str(val) -> str:
    if val is None:
        return ""
    s = str(val).replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"\s+", " ", s).strip()
    if len(s) > 400:
        s = s[:397] + "…"
    return s


def xlsx_bytes_to_pdf_bytes(content: bytes) -> bytes:
    """
    Строит PDF: каждый лист — заголовок + LongTable (страницы разбивает ReportLab).
    """
    from openpyxl import load_workbook
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    from xml.sax.saxutils import escape

    font = _ensure_ui_font()
    max_rows = max(100, int(getattr(config, "XLSX_PDF_MAX_ROWS_PER_SHEET", 6000) or 6000))
    max_cols = max(4, min(64, int(getattr(config, "XLSX_PDF_MAX_COLS", 32) or 32)))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )
    title_style = ParagraphStyle(
        name="SheetTitle",
        fontName=font,
        fontSize=11,
        leading=14,
    )
    elements: List = []

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            elements.append(Paragraph(f"<b>{escape(str(sheet_name)[:500])}</b>", title_style))
            elements.append(Spacer(1, 4))

            data_rows: List[List[str]] = []
            truncated = False
            for ri, row in enumerate(ws.iter_rows(values_only=True)):
                if ri >= max_rows:
                    truncated = True
                    break
                cells = list(row) if row else []
                line = [_cell_str(cells[j]) if j < len(cells) else "" for j in range(max_cols)]
                if len(cells) > max_cols:
                    line[-1] = "…"
                if any(line):
                    data_rows.append(line)

            if not data_rows:
                elements.append(Paragraph("<i>(пустой лист)</i>", title_style))
                elements.append(Spacer(1, 8))
                continue

            t = LongTable(data_rows, repeatRows=1, hAlign="LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, -1), font),
                        ("FONTSIZE", (0, 0), (-1, -1), 6),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E8E8")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.2, colors.grey),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFAFA")]),
                    ]
                )
            )
            elements.append(t)
            if truncated:
                elements.append(Spacer(1, 4))
                elements.append(
                    Paragraph(
                        f"<i>Обрезано: более {max_rows} строк на лист</i>",
                        title_style,
                    )
                )
            elements.append(Spacer(1, 10))
    finally:
        wb.close()

    if not elements:
        raise ValueError("Нет данных для PDF")

    doc.build(elements)
    pdf = buf.getvalue()
    if len(pdf) < 200:
        raise ValueError("PDF пустой или слишком мал")
    logger.info("XLSX→PDF: размер PDF %s байт", len(pdf))
    return pdf
