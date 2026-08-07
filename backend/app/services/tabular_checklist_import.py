"""Parse diligence checklists (txt/md/csv/docx) into column question lines."""
from __future__ import annotations

import csv
import io
import re
from typing import List, Optional, Tuple

_NUMBERED_RE = re.compile(
    r"^\s*(?:(?:\d+|[A-Za-zА-Яа-я])[\.\)\:]|[-*•–—])\s+(?P<body>.+?)\s*$"
)
_SUPPORTED_SUFFIXES = (".txt", ".md", ".csv", ".docx")


def supported_checklist_extensions() -> Tuple[str, ...]:
    return _SUPPORTED_SUFFIXES


def _clean_line(line: str) -> str:
    text = (line or "").replace("\u00a0", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def _extract_item(line: str) -> Optional[str]:
    cleaned = _clean_line(line)
    if not cleaned:
        return None
    if cleaned.startswith("#"):
        cleaned = cleaned.lstrip("#").strip()
        if not cleaned:
            return None
    match = _NUMBERED_RE.match(cleaned)
    if match:
        body = match.group("body").strip()
        return body or None
    # Plain sentence/question line (no numbering) — keep if looks like a checklist item.
    if len(cleaned) >= 3 and not cleaned.lower().startswith(("table of", "содержание")):
        return cleaned
    return None


def parse_checklist_text(text: str) -> List[str]:
    items: List[str] = []
    seen = set()
    for raw_line in (text or "").splitlines():
        item = _extract_item(raw_line)
        if not item:
            continue
        key = item.lower()
        if key in seen:
            continue
        seen.add(key)
        items.append(item)
    return items


def parse_checklist_csv(data: bytes) -> List[str]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    items: List[str] = []
    seen = set()
    for row in reader:
        if not row:
            continue
        # Prefer first non-empty cell; skip header-ish first row labels.
        cell = next((c.strip() for c in row if str(c).strip()), "")
        if not cell:
            continue
        if cell.lower() in {"question", "вопрос", "item", "пункт", "column", "колонка"}:
            continue
        item = _extract_item(cell) or cell
        key = item.lower()
        if key in seen:
            continue
        seen.add(key)
        items.append(item)
    return items


def parse_checklist_docx(data: bytes) -> List[str]:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines: List[str] = []
    for paragraph in document.paragraphs:
        text = _clean_line(paragraph.text)
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = _clean_line(cell.text)
                if text:
                    lines.append(text)
    return parse_checklist_text("\n".join(lines))


def parse_checklist_file(*, filename: str, data: bytes) -> List[str]:
    name = (filename or "").strip().lower()
    if not data:
        raise ValueError("Checklist file is empty")
    if name.endswith(".docx"):
        items = parse_checklist_docx(data)
    elif name.endswith(".csv"):
        items = parse_checklist_csv(data)
    elif name.endswith((".txt", ".md")):
        items = parse_checklist_text(data.decode("utf-8-sig", errors="replace"))
    else:
        allowed = ", ".join(_SUPPORTED_SUFFIXES)
        raise ValueError(f"Unsupported checklist format. Allowed: {allowed}")
    if not items:
        raise ValueError("No checklist items found in file")
    if len(items) > 80:
        raise ValueError("Checklist is too large (max 80 items)")
    return items


def questions_text_from_items(items: List[str]) -> str:
    return "\n".join(f"{idx}. {item}" for idx, item in enumerate(items, start=1))
